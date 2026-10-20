"""Cloudfone daily executive DOCX report generator with embedded visual charts.

Generates a publication-grade Word (.docx) document summarizing daily call center
performance, containing:
  1. Header with corporate identity and document reference.
  2. Executive KPI scorecard cards (Total calls, Service rate %, Missed, AHT).
  3. AI Insights & Highlights callout box (distilled from LLM daily notes).
  4. Visual charts & illustrations (Hourly distribution stacked bar, Donut rate).
  5. Detailed answered calls log with AI transcripts and summaries.
  6. Callback action tracking for missed calls.
  7. Recommended action plan and multi-party sign-off block.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Sequence

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import _Cell
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from app.services.cloudfone import (
    CallInsight,
    CallRecord,
    DailyCallStats,
    _fmt_duration,
    _hhmm,
)

# Brand Color Palette
COLOR_PRIMARY = RGBColor(27, 54, 93)  # Navy #1B365D
COLOR_SECONDARY = RGBColor(51, 102, 153)  # Slate Blue #336699
COLOR_DARK = RGBColor(30, 41, 59)  # Slate Dark #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)  # Cool Grey #64748B
COLOR_SUCCESS = RGBColor(22, 163, 74)  # Emerald Green #16A34A
COLOR_DANGER = RGBColor(220, 38, 38)  # Crimson Red #DC2626

HEX_HEADER_BG = "1B365D"
HEX_LIGHT_ROW = "F8FAFC"
HEX_CALLOUT_BG = "F0F4F8"
HEX_BORDER = "CBD5E1"
HEX_BORDER_LIGHT = "E2E8F0"
HEX_ACCENT_RED = "DC2626"
HEX_ACCENT_GREEN = "16A34A"


def _set_cell_margins(
    cell: _Cell,
    top: int = 100,
    bottom: int = 100,
    left: int = 120,
    right: int = 120,
) -> None:
    """Set inner cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_cell_shading(cell: _Cell, color_hex: str) -> None:
    """Set background color for a cell."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_borders(
    cell: _Cell,
    top: dict[str, str] | None = None,
    bottom: dict[str, str] | None = None,
    left: dict[str, str] | None = None,
    right: dict[str, str] | None = None,
) -> None:
    """Set explicit border styles per edge on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    borders = {"top": top, "bottom": bottom, "left": left, "right": right}
    for b_name, b_val in borders.items():
        if b_val:
            val_attr = b_val.get("val", "single")
            sz_attr = b_val.get("sz", "4")
            col_attr = b_val.get("color", "CBD5E1")
            b_elm = parse_xml(
                f'<w:{b_name} {nsdecls("w")} w:val="{val_attr}" w:sz="{sz_attr}" w:space="0" w:color="{col_attr}"/>'
            )
            tcBorders.append(b_elm)
        else:
            b_elm = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_elm)
    tcPr.append(tcBorders)


def _build_hourly_chart(calls: Sequence[CallRecord]) -> io.BytesIO:
    """Stacked bar chart showing answered vs missed calls across 4 daily slots."""
    slots = {
        "00:00 - 08:00\n(Sáng sớm)": {"ans": 0, "mis": 0},
        "08:00 - 12:00\n(Sáng)": {"ans": 0, "mis": 0},
        "12:00 - 18:00\n(Chiều)": {"ans": 0, "mis": 0},
        "18:00 - 24:00\n(Tối)": {"ans": 0, "mis": 0},
    }
    slot_keys = list(slots.keys())
    for call in calls:
        time_str = _hhmm(call.started_at)
        hour = int(time_str.split(":")[0]) if ":" in time_str else 9
        if hour < 8:
            key = slot_keys[0]
        elif hour < 12:
            key = slot_keys[1]
        elif hour < 18:
            key = slot_keys[2]
        else:
            key = slot_keys[3]

        if call.status == "ANSWERED":
            slots[key]["ans"] += 1
        else:
            slots[key]["mis"] += 1

    labels = list(slots.keys())
    ans_vals = [slots[k]["ans"] for k in labels]
    mis_vals = [slots[k]["mis"] for k in labels]
    totals = [a + m for a, m in zip(ans_vals, mis_vals)]

    fig, ax = plt.subplots(figsize=(6.4, 2.5), dpi=180)
    width = 0.44
    ax.bar(labels, ans_vals, label="Đã trả lời", color="#16a34a", width=width)
    ax.bar(labels, mis_vals, bottom=ans_vals, label="Cuộc gọi nhỡ", color="#dc2626", width=width)

    max_val = max(totals) if totals else 1
    ax.set_ylim(0, max(max_val + 2, 4))

    for idx, tot in enumerate(totals):
        if tot > 0:
            ax.text(
                idx,
                tot + 0.25,
                f"{tot} cuộc",
                ha="center",
                va="bottom",
                fontsize=8.5,
                fontweight="bold",
                color="#1e293b",
            )

    ax.set_title(
        "Phân bổ lưu lượng cuộc gọi theo khung giờ trong ngày",
        fontsize=10.5,
        fontweight="bold",
        color="#1e3a8a",
        pad=10,
    )
    ax.set_ylabel("Số lượng cuộc", fontsize=8.5, color="#64748b")
    ax.tick_params(axis="x", labelsize=8.5, colors="#334155")
    ax.tick_params(axis="y", labelsize=8, colors="#64748b")
    ax.grid(axis="y", linestyle="--", alpha=0.4, color="#cbd5e1")
    ax.set_axisbelow(True)
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)
    for spine in ["left", "bottom"]:
        ax.spines[spine].set_color("#cbd5e1")

    ax.legend(loc="upper right", frameon=True, fontsize=8, facecolor="#f8fafc", edgecolor="#e2e8f0")
    plt.tight_layout()
    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=200)
    plt.close(fig)
    buf.seek(0)
    return buf


def _build_summary_donut_chart(stats: DailyCallStats, calls: Sequence[CallRecord]) -> io.BytesIO:
    """Side-by-side visual: donut chart of service rate and direction distribution."""
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(6.4, 2.3), dpi=180)

    # Subplot 1: Donut chart
    ans_pct = stats.answer_rate_pct
    mis_pct = round(100.0 - ans_pct, 1) if stats.total_calls > 0 else 0.0

    if stats.total_calls > 0:
        sizes = [stats.answered_calls, stats.missed_calls]
        labels = [f"Trả lời\n({ans_pct:.1f}%)", f"Nhỡ\n({mis_pct:.1f}%)"]
        colors = ["#16a34a", "#dc2626"]
        pie_data = [(s, lbl, c) for s, lbl, c in zip(sizes, labels, colors) if s > 0]
        p_sizes = [x[0] for x in pie_data]
        p_labels = [x[1] for x in pie_data]
        p_colors = [x[2] for x in pie_data]

        ax1.pie(
            p_sizes,
            labels=p_labels,
            colors=p_colors,
            startangle=90,
            textprops=dict(color="#1e293b", fontsize=8, weight="bold"),
            wedgeprops=dict(width=0.42, edgecolor="white", linewidth=2),
        )
        ax1.text(
            0,
            0,
            f"{stats.total_calls}\ncuộc",
            ha="center",
            va="center",
            fontsize=9,
            fontweight="bold",
            color="#1e3a8a",
        )
    else:
        ax1.text(0.5, 0.5, "Không có cuộc gọi", ha="center", va="center", color="#64748b")
        ax1.axis("off")
    ax1.set_title(
        f"Tỷ lệ tiếp nhận ({ans_pct:.1f}%)",
        fontsize=9.5,
        fontweight="bold",
        color="#1e3a8a",
        pad=6,
    )

    # Subplot 2: Inbound vs Outbound
    inbound_count = sum(1 for c in calls if c.direction.lower() in ("inbound", "in", "gọi vào"))
    outbound_count = sum(1 for c in calls if c.direction.lower() in ("outbound", "out", "gọi ra"))
    if inbound_count == 0 and outbound_count == 0:
        inbound_count = stats.total_calls

    cats = ["Gọi vào (Inbound)", "Gọi ra (Outbound)"]
    vals = [inbound_count, outbound_count]
    bars = ax2.barh(cats, vals, color=["#2563eb", "#0d9488"], height=0.42)
    max_val = max(vals) if vals else 1
    ax2.set_xlim(0, max(max_val + 2, 4))
    for bar in bars:
        w = bar.get_width()
        ax2.text(
            w + 0.15,
            bar.get_y() + bar.get_height() / 2,
            f"{int(w)} cuộc",
            va="center",
            ha="left",
            fontsize=8,
            fontweight="bold",
            color="#334155",
        )
    ax2.set_title("Cơ cấu luồng cuộc gọi", fontsize=9.5, fontweight="bold", color="#1e3a8a", pad=6)
    ax2.tick_params(axis="y", labelsize=8, colors="#334155")
    ax2.tick_params(axis="x", labelsize=7.5, colors="#64748b")
    ax2.grid(axis="x", linestyle="--", alpha=0.4, color="#cbd5e1")
    ax2.set_axisbelow(True)
    for spine in ["top", "right"]:
        ax2.spines[spine].set_visible(False)
    for spine in ["left", "bottom"]:
        ax2.spines[spine].set_color("#cbd5e1")

    plt.tight_layout()
    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=200)
    plt.close(fig)
    buf.seek(0)
    return buf


def generate_daily_call_report_docx(
    stats: DailyCallStats,
    calls: Sequence[CallRecord],
    insights: Sequence[CallInsight],
    daily_notes: str | None,
    day: date,
) -> bytes:
    """Generate a complete daily call report DOCX document and return raw bytes."""
    doc = docx.Document()

    # Page setup - Standard A4 with 0.75 inch (54 pt) margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Styles
    normal_style = doc.styles["Normal"]
    normal_font = getattr(normal_style, "font", None)
    if normal_font is not None:
        normal_font.name = "Calibri"
        normal_font.size = Pt(10)
        normal_font.color.rgb = COLOR_DARK

    def add_sec_heading(title: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = COLOR_PRIMARY
        # subtle underline accent
        bdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="{HEX_BORDER}"/></w:pBdr>'
        )
        p._p.get_or_add_pPr().append(bdr)

    # 1. Corporate Header (2-column borderless table)
    hdr_tbl = doc.add_table(rows=1, cols=2)
    hdr_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = hdr_tbl.rows[0].cells[0]
    c_right = hdr_tbl.rows[0].cells[1]
    c_left.width = Inches(4.5)
    c_right.width = Inches(2.27)
    _set_cell_margins(c_left, top=0, bottom=60, left=0, right=0)
    _set_cell_margins(c_right, top=0, bottom=60, left=0, right=0)
    _set_cell_borders(c_left)
    _set_cell_borders(c_right)

    p_l = c_left.paragraphs[0]
    p_l.paragraph_format.space_before = Pt(0)
    p_l.paragraph_format.space_after = Pt(0)
    r_c = p_l.add_run("CÔNG TY TNHH VẬN TẢI & LOGISTICS TING TING\n")
    r_c.bold = True
    r_c.font.size = Pt(10)
    r_c.font.color.rgb = COLOR_PRIMARY
    r_sub = p_l.add_run("HỆ THỐNG GIÁM SÁT TỔNG ĐÀI CLOUDFONE & AI INSIGHTS")
    r_sub.font.size = Pt(8.5)
    r_sub.font.color.rgb = COLOR_MUTED

    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.space_before = Pt(0)
    p_r.paragraph_format.space_after = Pt(0)
    r_ref = p_r.add_run(
        f"Số: BC-CF/{day.strftime('%Y%m%d')}\nNgày báo cáo: {day.strftime('%d/%m/%Y')}"
    )
    r_ref.font.size = Pt(8.5)
    r_ref.font.color.rgb = COLOR_MUTED

    # Header Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(12)
    div_xml = parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="{HEX_HEADER_BG}"/></w:pBdr>'
    )
    p_div._p.get_or_add_pPr().append(div_xml)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(3)
    r_t = p_title.add_run("BÁO CÁO VẬN HÀNH TỔNG ĐÀI & PHÂN TÍCH CUỘC GỌI")
    r_t.bold = True
    r_t.font.size = Pt(15)
    r_t.font.color.rgb = COLOR_PRIMARY

    p_subt = doc.add_paragraph()
    p_subt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subt.paragraph_format.space_before = Pt(0)
    p_subt.paragraph_format.space_after = Pt(14)
    r_st = p_subt.add_run(
        f"Ngày theo dõi: {day.strftime('%d/%m/%Y')}   ·   Đầu số Hotline: 1800 7228   ·   Đơn vị: Ban Vận Hành"
    )
    r_st.italic = True
    r_st.font.size = Pt(9.5)
    r_st.font.color.rgb = COLOR_SECONDARY

    # 2. Executive KPI Scorecard (1x4 Table)
    kpi_tbl = doc.add_table(rows=1, cols=4)
    kpi_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    ans_pct = round(stats.answer_rate_pct, 1)
    mis_pct = round(100.0 - ans_pct, 1) if stats.total_calls > 0 else 0.0

    cards_data = [
        ("TỔNG CUỘC GỌI", str(stats.total_calls), "Toàn bộ lưu lượng", COLOR_PRIMARY),
        (
            "TỶ LỆ PHỤC VỤ",
            f"{ans_pct:.1f}%",
            f"{stats.answered_calls}/{stats.total_calls} đã nghe (SLA ≥ 80%)",
            COLOR_SUCCESS if ans_pct >= 80 else COLOR_DANGER,
        ),
        (
            "CUỘC BỊ NHỠ",
            str(stats.missed_calls),
            f"{mis_pct:.1f}% tổng cuộc",
            COLOR_DANGER if stats.missed_calls > 0 else COLOR_MUTED,
        ),
        (
            "THỜI LƯỢNG TB (AHT)",
            _fmt_duration(stats.avg_talk_seconds),
            "Mục tiêu SLA < 03:00",
            COLOR_PRIMARY,
        ),
    ]

    card_width = Inches(6.77 / 4.0)
    for idx, (label, val, sub, val_col) in enumerate(cards_data):
        cell = kpi_tbl.rows[0].cells[idx]
        cell.width = card_width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        _set_cell_shading(cell, HEX_LIGHT_ROW)
        _set_cell_borders(
            cell,
            top={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
            bottom={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
            left={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
            right={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
        )
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)

        rl = cp.add_run(f"{label}\n")
        rl.bold = True
        rl.font.size = Pt(7.5)
        rl.font.color.rgb = COLOR_MUTED

        rv = cp.add_run(f"{val}\n")
        rv.bold = True
        rv.font.size = Pt(17)
        rv.font.color.rgb = val_col

        rs = cp.add_run(sub)
        rs.font.size = Pt(7.5)
        rs.font.color.rgb = COLOR_MUTED

    # 3. Highlights & AI Insights Callout Box
    add_sec_heading("1. ĐIỂM TIN VẬN HÀNH & PHÁT HIỆN TRỌNG TÂM (AI INSIGHTS)")

    co_tbl = doc.add_table(rows=1, cols=1)
    co_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    co_cell = co_tbl.rows[0].cells[0]
    co_cell.width = Inches(6.77)
    _set_cell_margins(co_cell, top=120, bottom=120, left=160, right=160)
    _set_cell_shading(co_cell, HEX_CALLOUT_BG)
    _set_cell_borders(
        co_cell,
        left={"val": "single", "sz": "24", "color": HEX_HEADER_BG},
        top={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
        bottom={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
        right={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
    )

    co_p = co_cell.paragraphs[0]
    co_p.paragraph_format.space_before = Pt(2)
    co_p.paragraph_format.space_after = Pt(2)
    co_p.paragraph_format.line_spacing = 1.15
    r_cot = co_p.add_run("TÓM TẮT TÌNH HÌNH & NHẬN ĐỊNH TỰ ĐỘNG QUA AI:\n")
    r_cot.bold = True
    r_cot.font.size = Pt(9.5)
    r_cot.font.color.rgb = COLOR_PRIMARY

    if daily_notes:
        clean_notes = daily_notes.strip()
        r_cob = co_p.add_run(clean_notes)
        r_cob.font.size = Pt(9)
        r_cob.font.color.rgb = COLOR_DARK
    else:
        summary_text = (
            f"• Trong ngày {day.strftime('%d/%m/%Y')}, tổng đài đã ghi nhận {stats.total_calls} cuộc gọi.\n"
            f"• Tỷ lệ phục vụ đạt {ans_pct:.1f}% ({stats.answered_calls} cuộc trả lời thành công, {stats.missed_calls} cuộc nhỡ).\n"
            f"• Thời lượng đàm thoại bình quân (AHT) là {_fmt_duration(stats.avg_talk_seconds)} mỗi cuộc tiếp nhận."
        )
        r_cob = co_p.add_run(summary_text)
        r_cob.font.size = Pt(9)
        r_cob.font.color.rgb = COLOR_DARK

    # 4. Visual Graphs & Charts Section
    add_sec_heading("2. BIỂU ĐỒ TRỰC QUAN HÓA HOẠT ĐỘNG TỔNG ĐÀI")

    p_chart_intro = doc.add_paragraph()
    p_chart_intro.paragraph_format.space_before = Pt(2)
    p_chart_intro.paragraph_format.space_after = Pt(6)
    r_ci = p_chart_intro.add_run(
        "Hình 1 & 2: Biểu đồ phân bổ lưu lượng theo khung giờ và tỷ lệ tiếp nhận thực tế."
    )
    r_ci.italic = True
    r_ci.font.size = Pt(8.5)
    r_ci.font.color.rgb = COLOR_MUTED

    # Chart 1: Hourly distribution stacked bar
    hourly_stream = _build_hourly_chart(calls)
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_before = Pt(0)
    p_img1.paragraph_format.space_after = Pt(6)
    p_img1.add_run().add_picture(hourly_stream, width=Inches(6.6))

    # Chart 2: Summary donut & stream flow
    donut_stream = _build_summary_donut_chart(stats, calls)
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(0)
    p_img2.paragraph_format.space_after = Pt(10)
    p_img2.add_run().add_picture(donut_stream, width=Inches(6.6))

    # 5. Answered Calls Detail Table
    add_sec_heading("3. CHI TIẾT CÁC CUỘC GỌI ĐÃ PHỤC VỤ & PHÂN TÍCH AI")

    sorted_insights = sorted(insights, key=lambda item: item.transcript.started_at)
    if sorted_insights:
        c_widths = [Inches(0.85), Inches(1.1), Inches(1.0), Inches(0.85), Inches(2.97)]
        c_headers = [
            "Thời gian",
            "Số gọi",
            "Hướng / Nhánh",
            "Thời lượng",
            "Tóm tắt nội dung đàm thoại (AI)",
        ]

        tbl_ans = doc.add_table(rows=len(sorted_insights) + 1, cols=5)
        tbl_ans.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Table Header
        h_cells = tbl_ans.rows[0].cells
        for col_i, h_txt in enumerate(c_headers):
            h_cells[col_i].width = c_widths[col_i]
            _set_cell_margins(h_cells[col_i], top=80, bottom=80, left=80, right=80)
            _set_cell_shading(h_cells[col_i], HEX_HEADER_BG)
            _set_cell_borders(
                h_cells[col_i],
                top={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
                bottom={"val": "single", "sz": "12", "color": HEX_HEADER_BG},
            )
            hp = h_cells[col_i].paragraphs[0]
            hp.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if col_i in [0, 3]
                else (WD_ALIGN_PARAGRAPH.LEFT if col_i in [1, 2, 4] else WD_ALIGN_PARAGRAPH.LEFT)
            )
            hp.paragraph_format.space_before = Pt(0)
            hp.paragraph_format.space_after = Pt(0)
            hr = hp.add_run(h_txt)
            hr.bold = True
            hr.font.size = Pt(8.5)
            hr.font.color.rgb = RGBColor(255, 255, 255)

        # Table Rows
        for r_i, ins in enumerate(sorted_insights, start=1):
            call = ins.transcript
            r_cells = tbl_ans.rows[r_i].cells
            row_bg = HEX_LIGHT_ROW if r_i % 2 == 1 else "FFFFFF"

            time_val = _hhmm(call.started_at) or call.started_at[:16].replace("T", " ")
            caller_val = call.caller or "Không xác định"
            dir_agent = f"{call.direction or 'Inbound'} · {call.agent or 'N/A'}"
            dur_val = _fmt_duration(call.talk_seconds)
            summary_val = ins.summary or (
                call.transcript[:150] + "..." if call.transcript else "(Transcript trống)"
            )

            row_data = [time_val, caller_val, dir_agent, dur_val, summary_val]
            for col_i, text_val in enumerate(row_data):
                r_cells[col_i].width = c_widths[col_i]
                _set_cell_margins(r_cells[col_i], top=60, bottom=60, left=80, right=80)
                _set_cell_shading(r_cells[col_i], row_bg)
                _set_cell_borders(
                    r_cells[col_i],
                    top={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
                    bottom={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
                )
                rp = r_cells[col_i].paragraphs[0]
                rp.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if col_i in [0, 3]
                    else (
                        WD_ALIGN_PARAGRAPH.LEFT if col_i in [1, 2, 4] else WD_ALIGN_PARAGRAPH.LEFT
                    )
                )
                rp.paragraph_format.space_before = Pt(0)
                rp.paragraph_format.space_after = Pt(0)
                rr = rp.add_run(text_val)
                rr.font.size = Pt(8.5)
                if col_i == 1:
                    rr.bold = True
    else:
        p_no_ans = doc.add_paragraph()
        p_no_ans.paragraph_format.space_before = Pt(4)
        p_no_ans.paragraph_format.space_after = Pt(6)
        r_na = p_no_ans.add_run("Không có cuộc gọi nào được tiếp nhận trong ngày.")
        r_na.italic = True
        r_na.font.size = Pt(9)
        r_na.font.color.rgb = COLOR_MUTED

    # 6. Missed Calls & Callback Action Tracking
    add_sec_heading("4. THEO DÕI CUỘC GỌI NHỠ & KẾ HOẠCH GỌI LẠI (CALLBACK ACTION)")

    if stats.missed:
        p_mis_warn = doc.add_paragraph()
        p_mis_warn.paragraph_format.space_before = Pt(2)
        p_mis_warn.paragraph_format.space_after = Pt(6)
        r_mw = p_mis_warn.add_run(
            f"⚠️ Ghi nhận {len(stats.missed)} cuộc gọi nhỡ. Đề nghị bộ phận CSKH liên hệ lại trước 09h30 sáng:"
        )
        r_mw.bold = True
        r_mw.font.size = Pt(9)
        r_mw.font.color.rgb = COLOR_DANGER

        m_widths = [Inches(1.2), Inches(1.8), Inches(1.6), Inches(2.17)]
        m_headers = ["Thời gian phát sinh", "Số thuê bao", "Trạng thái", "Hành động CSKH đề xuất"]

        tbl_mis = doc.add_table(rows=len(stats.missed) + 1, cols=4)
        tbl_mis.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header
        mh_cells = tbl_mis.rows[0].cells
        for col_i, h_txt in enumerate(m_headers):
            mh_cells[col_i].width = m_widths[col_i]
            _set_cell_margins(mh_cells[col_i], top=80, bottom=80, left=80, right=80)
            _set_cell_shading(mh_cells[col_i], HEX_ACCENT_RED)
            _set_cell_borders(
                mh_cells[col_i],
                top={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
                bottom={"val": "single", "sz": "12", "color": HEX_ACCENT_RED},
            )
            mp = mh_cells[col_i].paragraphs[0]
            mp.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_i in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            mp.paragraph_format.space_before = Pt(0)
            mp.paragraph_format.space_after = Pt(0)
            mr = mp.add_run(h_txt)
            mr.bold = True
            mr.font.size = Pt(8.5)
            mr.font.color.rgb = RGBColor(255, 255, 255)

        # Rows
        for r_i, (started_at, caller) in enumerate(stats.missed, start=1):
            mr_cells = tbl_mis.rows[r_i].cells
            row_bg = HEX_LIGHT_ROW if r_i % 2 == 1 else "FFFFFF"
            time_txt = _hhmm(started_at) or started_at[:16].replace("T", " ")
            m_data = [
                time_txt,
                caller or "Không hiển thị",
                "Cuộc gọi nhỡ (NO ANSWER)",
                "Cần gọi lại hỗ trợ / ghi nhận nhu cầu",
            ]
            for col_i, txt in enumerate(m_data):
                mr_cells[col_i].width = m_widths[col_i]
                _set_cell_margins(mr_cells[col_i], top=60, bottom=60, left=80, right=80)
                _set_cell_shading(mr_cells[col_i], row_bg)
                _set_cell_borders(
                    mr_cells[col_i],
                    top={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
                    bottom={"val": "single", "sz": "4", "color": HEX_BORDER_LIGHT},
                )
                mp = mr_cells[col_i].paragraphs[0]
                mp.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if col_i in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
                )
                mp.paragraph_format.space_before = Pt(0)
                mp.paragraph_format.space_after = Pt(0)
                mr = mp.add_run(txt)
                mr.font.size = Pt(8.5)
                if col_i == 1:
                    mr.bold = True
    else:
        p_ok = doc.add_paragraph()
        p_ok.paragraph_format.space_before = Pt(2)
        p_ok.paragraph_format.space_after = Pt(6)
        r_ok = p_ok.add_run("✓ Không phát sinh cuộc gọi nhỡ trong ngày. Tỷ lệ tiếp nhận đạt 100%.")
        r_ok.bold = True
        r_ok.font.size = Pt(9)
        r_ok.font.color.rgb = COLOR_SUCCESS

    # 7. Action Plan & Sign-off Block
    add_sec_heading("5. ĐỀ XUẤT HÀNH ĐỘNG & XÁC NHẬN BÁO CÁO")

    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.space_before = Pt(2)
    p_rec.paragraph_format.space_after = Pt(8)
    p_rec.paragraph_format.line_spacing = 1.15
    p_rec.add_run(
        "1. Duy trì trực hotline liên tục trong giờ hành chính; rà soát cuộc gọi nhỡ để hoàn thành callback trong vòng 30 phút.\n"
        "2. Kiểm tra các trường hợp khiếu nại hoặc khách hàng gọi lặp lại để chuyển phòng ban chuyên trách xử lý dứt điểm.\n"
        "3. Đảm bảo chất lượng thoại và tốc độ nhấc máy dưới 3 hồi chuông theo chuẩn SLA dịch vụ."
    )

    sig_tbl = doc.add_table(rows=2, cols=3)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_widths = [Inches(2.25), Inches(2.25), Inches(2.27)]
    roles = [
        ("NGƯỜI LẬP BÁO CÁO", "Chuyên viên Hệ thống & AI Insights", "(Ký & ghi rõ họ tên)"),
        ("TRƯỞNG BỘ PHẬN VẬN HÀNH / CSKH", "Quản lý Vận hành Tổng đài", "(Ký & ghi rõ họ tên)"),
        ("BAN GIÁM ĐỐC", "Đại diện Lãnh đạo phê duyệt", "(Ký & ghi rõ họ tên)"),
    ]

    for col_i, (t1, t2, t3) in enumerate(roles):
        c_top = sig_tbl.rows[0].cells[col_i]
        c_top.width = s_widths[col_i]
        _set_cell_margins(c_top, top=60, bottom=40, left=40, right=40)
        _set_cell_borders(c_top)
        p1 = c_top.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)

        r1 = p1.add_run(f"{t1}\n")
        r1.bold = True
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_PRIMARY

        r2 = p1.add_run(f"{t2}\n")
        r2.italic = True
        r2.font.size = Pt(8)
        r2.font.color.rgb = COLOR_MUTED

        r3 = p1.add_run(t3)
        r3.font.size = Pt(7.5)
        r3.font.color.rgb = COLOR_MUTED

        c_bot = sig_tbl.rows[1].cells[col_i]
        c_bot.width = s_widths[col_i]
        _set_cell_margins(c_bot, top=200, bottom=40, left=40, right=40)
        _set_cell_borders(c_bot)
        p2 = c_bot.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(28)
        p2.paragraph_format.space_after = Pt(0)
        p2.add_run("")

    # Serialize to in-memory bytes
    out_stream = io.BytesIO()
    doc.save(out_stream)
    return out_stream.getvalue()
